"""Rebuild a two-column technical PDF as editable Word content.

This converter uses Poppler's PDF layout coordinates. It writes editable text
into Word table cells and inserts only the PDF's original raster images, never
whole-page screenshots.
"""

from __future__ import annotations

import argparse
import shutil
import subprocess
import tempfile
from collections import defaultdict
from pathlib import Path
from xml.etree import ElementTree as ET

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


def run(command: list[str], cwd: Path) -> None:
    subprocess.run(command, cwd=cwd, check=True, capture_output=True)


def set_cell_margins(cell, top=40, start=50, bottom=40, end=50) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def hide_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "nil")
        borders.append(element)
    tbl_pr.append(borders)


def clear_document(document: Document) -> None:
    body = document._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def parse_layout(xml_path: Path) -> tuple[list[dict], list[dict]]:
    root = ET.parse(xml_path).getroot()
    all_pages = []
    font_specs = {}
    for page in root.findall("page"):
        page_fonts = {
            spec.attrib["id"]: {
                "size": float(spec.attrib.get("size", "9")),
                "color": spec.attrib.get("color", "#000000"),
            }
            for spec in page.findall("fontspec")
        }
        rows = defaultdict(list)
        for text in page.findall("text"):
            value = "".join(text.itertext()).replace("\xa0", " ")
            if not value.strip():
                continue
            top = int(float(text.attrib.get("top", "0")))
            left = int(float(text.attrib.get("left", "0")))
            row_key = round(top / 3) * 3
            rows[row_key].append(
                {
                    "left": left,
                    "width": int(float(text.attrib.get("width", "0"))),
                    "value": value,
                    "font": page_fonts.get(text.attrib.get("font", ""), {"size": 9, "color": "#000000"}),
                }
            )
        lines = []
        for top, fragments in rows.items():
            fragments.sort(key=lambda item: item["left"])
            line = []
            last_end = None
            for fragment in fragments:
                if last_end is not None and fragment["left"] - last_end > 8:
                    line.append(" ")
                line.append(fragment["value"])
                last_end = fragment["left"] + fragment["width"]
            left = fragments[0]["left"]
            lines.append({"top": top, "left": left, "text": "".join(line).strip(), "font": fragments[0]["font"]})
        images = []
        for image in page.findall("image"):
            width = int(float(image.attrib.get("width", "0")))
            height = int(float(image.attrib.get("height", "0")))
            # Ignore decorative bullets/icons. They are represented as text in
            # most pages and make Word layout noisier when embedded separately.
            if width * height < 3_000:
                continue
            images.append(
                {
                    "top": int(float(image.attrib.get("top", "0"))),
                    "left": int(float(image.attrib.get("left", "0"))),
                    "width": width,
                    "height": height,
                    "src": image.attrib["src"],
                }
            )
        all_pages.append({"width": int(page.attrib["width"]), "lines": lines, "images": images})
        font_specs.update(page_fonts)
    return all_pages, font_specs


def add_text(cell, item: dict, page_width: int) -> None:
    paragraph = cell.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.line_spacing = 0.9
    run = paragraph.add_run(item["text"])
    size = max(6, min(16, item["font"]["size"] * 0.75))
    run.font.size = Pt(size)
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    if size >= 10.5:
        run.bold = True
    color = item["font"]["color"].lstrip("#")
    if len(color) == 6:
        run.font.color.rgb = RGBColor.from_string(color)


def add_image(cell, image: dict, base_dir: Path, page_width: int) -> None:
    path = base_dir / image["src"]
    if not path.is_file():
        return
    paragraph = cell.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(1)
    ratio = image["width"] / max(page_width, 1)
    width = min(3.55, max(0.7, ratio * 7.6))
    paragraph.add_run().add_picture(str(path), width=Inches(width))


def build_document(source: Path, template: Path, output: Path) -> None:
    with tempfile.TemporaryDirectory(prefix="pdf_word_layout_") as temp:
        work = Path(temp)
        staged = work / "source.pdf"
        shutil.copy2(source, staged)
        run(["pdftohtml", "-xml", "-enc", "UTF-8", "source.pdf", "layout.xml"], work)
        pages, _ = parse_layout(work / "layout.xml")

        document = Document(template)
        clear_document(document)
        section = document.sections[0]
        section.top_margin = Inches(0.35)
        section.bottom_margin = Inches(0.35)
        section.left_margin = Inches(0.35)
        section.right_margin = Inches(0.35)

        for page_index, page in enumerate(pages):
            table = document.add_table(rows=1, cols=2)
            table.autofit = False
            hide_table_borders(table)
            left_cell, right_cell = table.rows[0].cells
            for cell in (left_cell, right_cell):
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
                set_cell_margins(cell)
                cell.paragraphs[0]._element.getparent().remove(cell.paragraphs[0]._element)

            events = [dict(item, kind="text") for item in page["lines"]]
            events.extend(dict(item, kind="image") for item in page["images"])
            by_column = {0: [], 1: []}
            for event in events:
                center = event["left"] + event.get("width", 0) / 2
                by_column[0 if center < page["width"] / 2 else 1].append(event)
            for column, cell in ((0, left_cell), (1, right_cell)):
                for event in sorted(by_column[column], key=lambda item: (item["top"], item["left"])):
                    if event["kind"] == "text":
                        add_text(cell, event, page["width"])
                    else:
                        add_image(cell, event, work, page["width"])

            if page_index < len(pages) - 1:
                document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

        document.save(output)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("pdf", type=Path)
    parser.add_argument("--template", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    build_document(args.pdf.resolve(), args.template.resolve(), args.output.resolve())


if __name__ == "__main__":
    main()
